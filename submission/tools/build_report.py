"""Build the final report .docx from submission/report/final_report.md.

    python submission/tools/build_report.py

The Markdown file is the single source of truth for the report text. This
script supports only the small Markdown subset that file uses: a front-matter
block, ## / ### headings, paragraphs with **bold** and *italic*, "- " bullets,
"1. " numbered items, pipe tables, and figures written as
![Figure N. caption](path){width=12cm}. Numbered items are written as literal
text with a hanging indent, so separate lists never share one Word numbering
sequence.

The PDF is exported from the .docx by Microsoft Word with
submission/tools/export_pdf.ps1, so the two files always match.
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

SUBMISSION = Path(__file__).resolve().parents[1]
SOURCE = SUBMISSION / "report" / "final_report.md"
OUTPUT = SUBMISSION / "report" / "Tashkent_Transit_Walkability_Final_Report.docx"

INK = RGBColor(0x14, 0x1B, 0x24)
MUTED = RGBColor(0x55, 0x63, 0x73)
ACCENT = RGBColor(0x0E, 0x6F, 0x7E)      # darkened site metro teal, readable on white
HEADER_FILL = "0E3B45"
ZEBRA_FILL = "F1F5F7"
BODY_FONT = "Calibri"
MAX_IMAGE_PX = 2400

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
FIGURE = re.compile(r"^!\[(?P<cap>.+?)\]\((?P<path>[^)]+)\)(\{width=(?P<w>[\d.]+)cm\})?\s*$")
NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")


def set_cell_fill(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def bottom_border(paragraph, colour="0E6F7E", size=12) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), colour)
    borders.append(bottom)
    p_pr.append(borders)


def add_hyperlink(paragraph, url: str, text: str | None = None) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "0E6F7E")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(colour)
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text or url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


URL = re.compile(r"(https?://\S+?)([.,;)]?)(?=\s|$)")


def typographic(text: str) -> str:
    """Curly quotes and apostrophes for body text (URLs never pass through here)."""
    text = re.sub(r"(?<=\w)'(?=\w)", "\u2019", text)          # Tashkent's
    text = re.sub(r"(^|[\s(\[])'", "\\1\u2018", text)          # opening single
    text = text.replace("'", "\u2019")
    text = re.sub(r'(^|[\s(\[])"', "\\1\u201c", text)          # opening double
    return text.replace('"', "\u201d")


def add_inline(paragraph, text: str, size: float | None = None, colour: RGBColor | None = None,
               link_urls: bool = False) -> None:
    """Write **bold** / *italic* runs; optionally turn bare URLs into hyperlinks."""
    if not link_urls:
        text = typographic(text)
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        bold = chunk.startswith("**") and chunk.endswith("**")
        italic = not bold and chunk.startswith("*") and chunk.endswith("*")
        body = chunk[2:-2] if bold else chunk[1:-1] if italic else chunk
        pieces = URL.split(body) if link_urls else [body]
        if link_urls:
            # URL.split yields [text, url, trailing, text, url, trailing, ...]
            i = 0
            while i < len(pieces):
                if i % 3 == 1:
                    add_hyperlink(paragraph, pieces[i])
                elif pieces[i]:
                    run = paragraph.add_run(pieces[i])
                    run.bold, run.italic = bold, italic
                    if size:
                        run.font.size = Pt(size)
                    if colour is not None:
                        run.font.color.rgb = colour
                i += 1
            continue
        run = paragraph.add_run(body)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if colour is not None:
            run.font.color.rgb = colour


def parse_front_matter(text: str) -> tuple[dict[str, str], str]:
    if not text.startswith("---"):
        return {}, text
    _, block, rest = text.split("---", 2)
    meta = {}
    for line in block.strip().splitlines():
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip().strip('"')
    return meta, rest


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.2))
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, colour, before in (("Heading 1", 15, ACCENT, 16), ("Heading 2", 12, INK, 10)):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = colour
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet",):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)


def add_footer(doc: Document, short_title: str) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"{short_title}  ·  STATS 401 Final Report  ·  page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    num = p.add_run()
    num.font.size = Pt(8.5)
    num.font.color.rgb = MUTED
    add_field(num, "PAGE")


def add_title_block(doc: Document, meta: dict[str, str]) -> None:
    eyebrow = doc.add_paragraph()
    run = eyebrow.add_run(meta.get("course", "").upper())
    run.font.size, run.bold = Pt(9), True
    run.font.color.rgb = ACCENT
    eyebrow.paragraph_format.space_after = Pt(2)

    title = doc.add_paragraph()
    run = title.add_run(meta["title"])
    run.font.size, run.bold = Pt(24), True
    run.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(0)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run(meta.get("subtitle", ""))
    run.font.size = Pt(14)
    run.font.color.rgb = ACCENT
    subtitle.paragraph_format.space_after = Pt(6)

    authors = doc.add_paragraph()
    run = authors.add_run(meta.get("authors", ""))
    run.font.size, run.bold = Pt(10.5), True
    run = authors.add_run(f"   ·   {meta.get('date', '')}")
    run.font.size = Pt(10.5)
    run.font.color.rgb = MUTED
    authors.paragraph_format.space_after = Pt(2)

    links = doc.add_paragraph()
    for label, key in (("Live website: ", "site"), ("   Code and data: ", "repo")):
        run = links.add_run(label)
        run.font.size = Pt(9.5)
        run.font.color.rgb = MUTED
        add_hyperlink(links, meta[key])
    for r in links.runs:
        r.font.size = Pt(9.5)
    bottom_border(links)
    links.paragraph_format.space_after = Pt(10)


def column_widths(header: list[str]) -> list[float]:
    usable = 16.6
    key = [h.split(":")[0].strip().lower() for h in header]
    if key[:2] == ["source", "role in the project"]:
        return [3.4, 8.0, 5.2]
    if key[:1] == ["view"]:
        return [2.7, 4.0, 3.4, 6.5]
    return [usable / len(header)] * len(header)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    header = rows[0]
    table = doc.add_table(rows=len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = doc.styles["Table Grid"]
    set_cell_margins(table)
    widths = column_widths(header)
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.cell(r_index, c_index)
            cell.width = Cm(widths[c_index])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            # Keep the whole table on one page: every row but the last binds to the next.
            p.paragraph_format.keep_with_next = r_index < len(rows) - 1
            if r_index == 0:
                set_cell_fill(cell, HEADER_FILL)
                run = p.add_run(value)
                run.bold = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                if r_index % 2 == 0:
                    set_cell_fill(cell, ZEBRA_FILL)
                add_inline(p, value if c_index else f"**{value}**", size=8.5)
    # Light grey grid instead of black.
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "C9D3DC")
        borders.append(node)
    table._tbl.tblPr.append(borders)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def image_stream(path: Path) -> io.BytesIO:
    with Image.open(path) as im:
        im = im.convert("RGB")
        if im.width > MAX_IMAGE_PX:
            ratio = MAX_IMAGE_PX / im.width
            im = im.resize((MAX_IMAGE_PX, round(im.height * ratio)), Image.LANCZOS)
        buffer = io.BytesIO()
        im.save(buffer, format="PNG", optimize=True)
    buffer.seek(0)
    return buffer


def add_figure(doc: Document, caption: str, path: Path, width_cm: float) -> None:
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.paragraph_format.space_before = Pt(4)
    pic.paragraph_format.space_after = Pt(2)
    pic.add_run().add_picture(image_stream(path), width=Cm(width_cm))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    label, _, rest = typographic(caption).partition(". ")
    run = cap.add_run(label + ". ")
    run.bold, run.font.size = True, Pt(8.5)
    run.font.color.rgb = ACCENT
    run = cap.add_run(rest)
    run.italic, run.font.size = True, Pt(8.5)
    run.font.color.rgb = MUTED


def build() -> Path:
    meta, body = parse_front_matter(SOURCE.read_text(encoding="utf-8"))
    doc = Document()
    configure(doc)
    add_footer(doc, "Who Can Reach the Metro in Ten Minutes?")
    add_title_block(doc, meta)

    lines = body.strip().splitlines()
    in_references = False
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            in_references = heading.lower() == "references"
            doc.add_heading(heading, level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    block.append(cells)
                i += 1
            add_table(doc, block)
            continue
        elif (m := FIGURE.match(line)):
            add_figure(doc, m["cap"], (SOURCE.parent / m["path"]).resolve(), float(m["w"] or 16))
        elif line.startswith("- "):
            if in_references:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.first_line_indent = Cm(-0.8)
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, line[2:], size=9, link_urls=True)
            else:
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, line[2:])
        elif (m := NUMBERED.match(line)):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
            run = p.add_run(f"{m[1]}.\t")
            run.bold = True
            run.font.color.rgb = ACCENT
            add_inline(p, m[2])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    out = build()
    print(f"wrote {out.relative_to(SUBMISSION.parent)} ({out.stat().st_size / 1e6:.1f} MB)")
    sys.exit(0)
